from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from apps.core.datas import hoje
from apps.relatorios.consultas import montar_aba
from apps.relatorios.exportacao import para_excel, para_pdf
from tests.apoio import F, criar_cliente, montar_pedido_confirmado, texto_do_pdf

SET = F(date(2026, 9, 1), date(2026, 9, 30))
FIM_DO_ANO = date(2026, 12, 1)
EXCEL = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _planilhas(aba, f, dias=60):
    dados = para_excel(montar_aba(aba, f, FIM_DO_ANO, dias), f)
    return load_workbook(BytesIO(dados))


def test_excel_com_numeros_de_verdade(cenario):
    ws = _planilhas("vendas", SET).active
    linha = next(r for r in ws.iter_rows() if r[0].value == "set/26")
    liquido = linha[4]
    assert liquido.data_type == "n"
    assert Decimal(str(liquido.value)) == Decimal("325.00") and "R$" in liquido.number_format
    assert linha[6].number_format == "0.0%"
    assert "Vendas por período" in ws["A1"].value and "01/09/2026" in ws["A1"].value
    assert any(c.font.bold and c.value == "Líquido" for r in ws.iter_rows() for c in r)


def test_excel_uma_planilha_por_tabela_com_nome_curto(cenario):
    wb = _planilhas("clientes", SET)
    assert len(wb.sheetnames) == 2 and all(len(n) <= 31 for n in wb.sheetnames)
    assert wb.sheetnames[0] == "Ranking de clientes"


def test_excel_datas_e_filtros(cenario):
    wb = _planilhas(
        "clientes", F(date(2026, 9, 1), date(2026, 10, 31), funcionario_id=cenario.carla.pk)
    )
    ranking = wb.worksheets[0]
    assert "Funcionário: Carla Souza" in ranking["A1"].value
    datas = [c for r in ranking.iter_rows() for c in r if c.data_type == "d"]
    assert datas and all(c.number_format == "dd/mm/yyyy" for c in datas)


def test_excel_nunca_cria_formula_com_dado_de_cadastro(cenario, administrador):
    cli = criar_cliente('=HYPERLINK("http://x","clique")')
    montar_pedido_confirmado(administrador, cliente=cli, itens=[(cenario.ce285a, 1)])
    dia = hoje()
    wb = _planilhas("clientes", F(dia, dia))
    celulas = [
        c
        for ws in wb
        for r in ws.iter_rows()
        for c in r
        if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")
    ]
    assert celulas and all(c.data_type == "s" for c in celulas)


def test_excel_estoque_tem_o_valor_total(cenario):
    ws = _planilhas("estoque", SET)["Valor total em estoque"]
    assert ws["A3"].value == "Valor total em estoque"
    assert Decimal(str(ws["A4"].value)) == Decimal("7740.00")


def test_pdf_valido(cenario):
    dados = para_pdf(montar_aba("vendas", SET, FIM_DO_ANO), SET)
    texto = texto_do_pdf(dados)
    assert dados.startswith(b"%PDF-") and "Vendas por período" in texto and "R$ 325,00" in texto
    assert "01/09/2026 a 30/09/2026" in texto


@pytest.mark.parametrize(
    "aba", ["resumo", "vendas", "clientes", "produtos", "funcionarios", "estoque", "cancelamentos"]
)
def test_todas_as_abas_exportam(cenario, aba):
    rel = montar_aba(aba, SET, FIM_DO_ANO)
    assert para_pdf(rel, SET).startswith(b"%PDF-")
    assert load_workbook(BytesIO(para_excel(rel, SET))).sheetnames


def test_pdf_do_estoque_mostra_o_valor_total(cenario):
    texto = texto_do_pdf(para_pdf(montar_aba("estoque", SET, FIM_DO_ANO), SET))
    assert "R$ 7.740,00" in texto


def test_download_respeita_os_filtros(client_admin, cenario):
    r = client_admin.get("/relatorios/vendas/excel/?atalho=datas&inicio=2026-09-01&fim=2026-09-30")
    assert 'filename="relatorio-vendas-2026-09-01-a-2026-09-30.xlsx"' in r["Content-Disposition"]
    assert r["Content-Type"] == EXCEL
    assert "no-store" in r["Cache-Control"] and "private" in r["Cache-Control"]
    ws = load_workbook(BytesIO(r.content)).active
    assert any(c.value == "set/26" for row in ws.iter_rows() for c in row)


def test_download_do_pdf(client_admin, cenario):
    r = client_admin.get("/relatorios/vendas/pdf/?atalho=datas&inicio=2026-09-01&fim=2026-09-30")
    assert 'filename="relatorio-vendas-2026-09-01-a-2026-09-30.pdf"' in r["Content-Disposition"]
    assert r["Content-Type"] == "application/pdf" and r.content.startswith(b"%PDF-")
    assert "no-store" in r["Cache-Control"] and "private" in r["Cache-Control"]


def test_exportacao_do_cliente_le_os_dias(client_admin, cenario):
    r = client_admin.get(
        "/relatorios/clientes/excel/?atalho=datas&inicio=2026-09-01&fim=2026-09-30&dias=45"
    )
    assert (
        "Clientes sem comprar há mais de 45 dias"
        in load_workbook(BytesIO(r.content)).worksheets[1]["A1"].value
    )
    r = client_admin.get(
        "/relatorios/clientes/excel/?atalho=datas&inicio=2026-09-01&fim=2026-09-30&dias=abc"
    )
    assert "mais de 60 dias" in load_workbook(BytesIO(r.content)).worksheets[1]["A1"].value


def test_aba_ou_formato_desconhecido_da_404(client_admin):
    assert client_admin.get("/relatorios/nada/excel/").status_code == 404
    assert client_admin.get("/relatorios/vendas/csv/").status_code == 404


def test_falha_na_exportacao_avisa(client_admin, cenario, monkeypatch):
    monkeypatch.setattr("apps.relatorios.views.para_pdf", Mock(side_effect=RuntimeError))
    r = client_admin.get("/relatorios/vendas/pdf/?atalho=12m", follow=True)
    assert r.redirect_chain[-1][0].startswith("/relatorios/vendas/?atalho=12m")
    assert (
        "Não foi possível gerar o arquivo. Tente de novo; se continuar, avise o administrador."
        in r.content.decode()
    )


def test_links_de_exportacao_usam_a_rota_real(client_admin):
    url = "/relatorios/vendas/?atalho=3m"
    query = "atalho=3m&amp;agrupamento=mes"
    for html in (
        client_admin.get(url).content.decode(),
        client_admin.get(url, headers={"HX-Request": "true"}).content.decode(),
    ):
        assert f'<a href="/relatorios/vendas/excel/?{query}" class="btn-sec">⤓ Excel</a>' in html
        assert f'<a href="/relatorios/vendas/pdf/?{query}" class="btn-sec">⤓ PDF</a>' in html


def test_link_de_exportacao_da_aba_clientes_leva_os_dias(client_admin):
    html = client_admin.get("/relatorios/clientes/?atalho=3m&dias=45").content.decode()
    assert "/relatorios/clientes/excel/?atalho=3m&amp;agrupamento=mes&amp;dias=45" in html
