"""Exportação dos relatórios para Excel (openpyxl) e PDF (`DocumentoPDF`), na memória."""

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.contas.models import Usuario
from apps.core.formatacao import brl, data_br, inteiro_br, percentual
from apps.core.pdf import CORPO, FONTE, DocumentoPDF

from .consultas import Relatorio
from .periodos import NOMES_DO_AGRUPAMENTO, Filtros
from .tabelas import Coluna, Tabela

FORMATO_DO_EXCEL = {
    "dinheiro": '"R$" #,##0.00',
    "percentual": "0.0%",
    "data": "dd/mm/yyyy",
    "inteiro": "#,##0",
}
MAXIMO_NOME_DA_PLANILHA = 31
FOLGA_DA_CELULA = 3  # mm: o padding da célula do PDF (1 mm de cada lado) e um respiro
MINIMO_DO_TEXTO = 30  # mm
VAZIO_DA_TABELA = "Nada para mostrar com estes filtros."


def descricao_dos_filtros(f: Filtros) -> str:
    """'Período: 01/09/2026 a 30/09/2026 · Agrupado por mês · Funcionário: Todos'."""
    if f.funcionario_id is None:
        funcionario = "Todos"
    else:
        nome = Usuario.objects.filter(pk=f.funcionario_id).values_list("nome", flat=True).first()
        funcionario = nome or "não encontrado"
    return (
        f"Período: {data_br(f.inicio)} a {data_br(f.fim)}"
        f" · Agrupado por {NOMES_DO_AGRUPAMENTO[f.agrupamento]}"
        f" · Funcionário: {funcionario}"
    )


# ---------- Excel ----------


def _nome_da_planilha(titulo: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", titulo)[:MAXIMO_NOME_DA_PLANILHA].strip() or "Planilha"
    nome, n = base, 2
    while nome.lower() in usados:
        sufixo = f" ({n})"
        nome = base[: MAXIMO_NOME_DA_PLANILHA - len(sufixo)] + sufixo
        n += 1
    usados.add(nome.lower())
    return nome


def _gravar_celula(ws, linha: int, coluna: int, valor: object, tipo: str) -> None:
    celula = ws.cell(row=linha, column=coluna)
    if tipo == "texto" or valor is None:
        celula.value = "" if valor is None else str(valor)
        # Sempre texto: um nome de cliente começando com "=" nunca pode virar fórmula.
        celula.data_type = "s"
        return
    celula.value = valor
    celula.number_format = FORMATO_DO_EXCEL[tipo]


def _planilha(ws, rel: Relatorio, tabela: Tabela, filtros: str) -> None:
    ws.cell(row=1, column=1, value=f"{rel.titulo} · {tabela.titulo} · {filtros}").font = Font(
        bold=True
    )
    ws["A1"].data_type = "s"
    larguras = [len(c.rotulo) for c in tabela.colunas]
    for j, coluna in enumerate(tabela.colunas, start=1):
        ws.cell(row=3, column=j, value=coluna.rotulo).font = Font(bold=True)
    for i, linha in enumerate(tabela.linhas, start=4):
        for j, (coluna, valor) in enumerate(zip(tabela.colunas, linha, strict=True), start=1):
            _gravar_celula(ws, i, j, valor, coluna.tipo)
            larguras[j - 1] = max(larguras[j - 1], len(_texto(coluna, valor)))
    for j, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(largura + 2, 10), 60)
    ws.freeze_panes = "A4"


def para_excel(rel: Relatorio, f: Filtros) -> bytes:
    """Uma planilha por tabela: título e filtros na linha 1, cabeçalho em negrito."""
    filtros = descricao_dos_filtros(f)
    wb = Workbook()
    wb.remove(wb.active)
    usados: set[str] = set()
    for tabela in rel.tabelas:
        ws = wb.create_sheet(_nome_da_planilha(tabela.titulo, usados))
        _planilha(ws, rel, tabela, filtros)
    if not rel.tabelas:
        wb.create_sheet("Relatório")
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ---------- PDF ----------


def _texto(coluna: Coluna, valor: object) -> str:
    """O texto da célula como a tela mostra."""
    if coluna.tipo == "texto":
        return str(valor) if valor not in (None, "") else "—"
    if coluna.tipo == "dinheiro":
        return brl(valor)
    if coluna.tipo == "inteiro":
        return "—" if valor is None else inteiro_br(valor)
    if coluna.tipo == "percentual":
        return percentual(valor)
    return data_br(valor)


def _largura_do_texto(pdf: DocumentoPDF, texto: str, negrito: bool = False) -> float:
    pdf.set_font(FONTE, "B" if negrito else "", CORPO)
    return pdf.get_string_width(texto)


def _larguras(pdf: DocumentoPDF, colunas: list[Coluna], linhas: list[list[str]]) -> list[float]:
    """Larguras em mm, medidas na fonte do PDF.

    Colunas de número, dinheiro e data nunca ficam menores que o maior texto delas (o valor não
    quebra no meio); o resto da página vai para as colunas de texto, na proporção do que
    precisam. Se nem os números couberem, eles encolhem juntos, nunca abaixo do cabeçalho.
    """
    cabecalho = [_largura_do_texto(pdf, c.rotulo, negrito=True) + FOLGA_DA_CELULA for c in colunas]
    desejada = [
        max(
            [cabecalho[j]]
            + [_largura_do_texto(pdf, linha[j]) + FOLGA_DA_CELULA for linha in linhas]
        )
        for j in range(len(colunas))
    ]
    texto = [j for j, c in enumerate(colunas) if c.tipo == "texto"]
    numeros = [j for j, c in enumerate(colunas) if c.tipo != "texto"]
    disponivel = pdf.epw
    if not texto:
        return [d * disponivel / sum(desejada) for d in desejada]
    larguras = list(desejada)
    minimo_texto = {j: min(desejada[j], MINIMO_DO_TEXTO) for j in texto}
    sobra = disponivel - sum(desejada[j] for j in numeros)
    if sobra < sum(minimo_texto.values()):
        # Os números encolhem juntos para dar o mínimo ao texto, sem passar do cabeçalho.
        total_numeros = sum(desejada[j] for j in numeros)
        fator = max(disponivel - sum(minimo_texto.values()), 0) / total_numeros if numeros else 0
        for j in numeros:
            larguras[j] = max(desejada[j] * fator, cabecalho[j])
        # Se nem assim couber, o texto divide o que restou, na proporção do mínimo de cada um.
        resto = max(disponivel - sum(larguras[j] for j in numeros), 1.0)
        escala = min(1.0, resto / sum(minimo_texto.values()))
        for j in texto:
            larguras[j] = minimo_texto[j] * escala
        return larguras
    peso = sum(desejada[j] for j in texto)
    for j in texto:
        larguras[j] = sobra * desejada[j] / peso
    return larguras


def para_pdf(rel: Relatorio, f: Filtros) -> bytes:
    """O relatório em A4 paisagem: título, filtros e uma seção por tabela."""
    pdf = DocumentoPDF(rel.titulo, descricao_dos_filtros(f), orientacao="L")
    for tabela in rel.tabelas:
        pdf.secao(tabela.titulo)
        if not tabela.linhas:
            pdf.cell(0, 6, VAZIO_DA_TABELA, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(4)
            continue
        linhas = [
            [_texto(c, v) for c, v in zip(tabela.colunas, linha, strict=True)]
            for linha in tabela.linhas
        ]
        larguras = _larguras(pdf, tabela.colunas, linhas)
        pdf.tabela(
            [
                (c.rotulo, largura, "L" if c.tipo == "texto" else "R")
                for c, largura in zip(tabela.colunas, larguras, strict=True)
            ],
            linhas,
        )
        pdf.ln(4)
    return pdf.para_bytes()
